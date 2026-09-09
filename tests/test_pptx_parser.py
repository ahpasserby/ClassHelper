"""pptx parsing, on decks built here rather than on any real course material.

The AlternateContent case is the one that matters. python-pptx does not descend
into those wrappers, so the shapes inside are absent from `slide.shapes` with no
error raised -- the failure mode is a slide that quietly loses its body text.
Real decks hit it whenever a shape uses a post-2010 drawing feature.
"""

from __future__ import annotations

import copy

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from classhelper.model import BlockKind
from classhelper.parsers import UnsupportedFormat, parse
from classhelper.parsers.pptx_parser import parse as parse_pptx

_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _wrap_in_alternate_content(shape) -> None:
    """Reproduce what PowerPoint emits for a shape needing a compatibility
    fallback: <mc:AlternateContent><mc:Choice>{shape}</mc:Choice>
    <mc:Fallback>{degraded shape}</mc:Fallback></mc:AlternateContent>."""
    sp = shape._element
    parent = sp.getparent()
    at = parent.index(sp)

    ac = etree.SubElement(parent, f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC})
    choice = etree.SubElement(ac, f"{{{_MC}}}Choice")
    choice.set("Requires", "a14")
    fallback = etree.SubElement(ac, f"{{{_MC}}}Fallback")

    parent.remove(sp)
    choice.append(sp)
    fallback.append(copy.deepcopy(sp))
    parent.remove(ac)
    parent.insert(at, ac)


def _add_equation(paragraph, text: str) -> None:
    """Append an OMML equation to a paragraph, the way PowerPoint stores one:
    Cambria Math's Unicode alphanumerics inside <m:oMath>, not a text run."""
    math_alpha = {c: chr(0x1D400 + ord(c) - ord("A")) for c in
                  "ABCDEFGHIJKLMNOPQRSTUVWXYZ"}
    styled = "".join(math_alpha.get(c, c) for c in text)

    omath = etree.SubElement(paragraph._p, f"{{{_M}}}oMath", nsmap={"m": _M})
    run = etree.SubElement(omath, f"{{{_M}}}r")
    node = etree.SubElement(run, f"{{{_M}}}t")
    node.text = styled


def _add_run(paragraph, text: str) -> None:
    run = paragraph.add_run()
    run.text = text


@pytest.fixture
def deck_path(tmp_path):
    """A two-slide deck whose body text is hidden inside AlternateContent."""
    prs = Presentation()
    for i in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = f"Slide {i + 1}"

        body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        tf = body.text_frame
        tf.text = "The first sentence. The second sentence."
        tf.add_paragraph().text = "A second paragraph."
        _wrap_in_alternate_content(body)

        nav = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(0.4))
        nav.text_frame.text = "Course Navigation"

        code = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(0.5))
        run = code.text_frame.paragraphs[0].add_run()
        run.text = "SELECT * FROM student;"
        run.font.name = "Consolas"
        run.font.size = Pt(14)

    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return str(path)


def test_text_inside_alternate_content_is_recovered(deck_path):
    deck = parse_pptx(deck_path)
    texts = [b.text for p in deck.pages for b in p.blocks]
    assert "The first sentence. The second sentence." in texts
    assert "A second paragraph." in texts
    # A diagnostic, not a warning: the reader never shows it, because
    # "mc:AlternateContent" means nothing to someone reading a lecture.
    assert any("AlternateContent" in n for n in deck.notes)
    assert deck.warnings == []


def test_the_fallback_branch_is_not_duplicated(deck_path):
    """Choice and Fallback describe the same shape. Emitting both would double
    every sentence, and the user would pay to translate each one twice."""
    deck = parse_pptx(deck_path)
    # Scoped to one slide: the fixture puts the same text on both.
    texts = [b.text for b in deck.pages[0].blocks]
    assert texts.count("A second paragraph.") == 1


def test_each_paragraph_becomes_its_own_block(deck_path):
    deck = parse_pptx(deck_path)
    page = deck.pages[0]
    assert "The first sentence. The second sentence." in [b.text for b in page.blocks]
    assert "A second paragraph." in [b.text for b in page.blocks]


def test_title_placeholder_is_detected(deck_path):
    from classhelper.classify import classify

    deck = parse_pptx(deck_path)
    classify(deck)
    assert deck.pages[0].title == "Slide 1"


def test_monospace_run_is_classified_as_code(deck_path):
    from classhelper.classify import classify

    deck = parse_pptx(deck_path)
    classify(deck)
    code = [b for p in deck.pages for b in p.blocks if b.kind is BlockKind.CODE]
    assert [b.text for b in code] == ["SELECT * FROM student;"] * 2


def test_blocks_come_back_in_reading_order(deck_path):
    deck = parse_pptx(deck_path)
    ys = [b.box.y for b in deck.pages[0].blocks]
    assert ys == sorted(ys), "blocks must be top-to-bottom, not XML order"


def test_equation_text_is_not_dropped(tmp_path):
    """Regression: python-pptx's `.text` skips OMML, so a sentence containing an
    equation arrived with a hole in it and the translator wrote around it."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    para = box.text_frame.paragraphs[0]
    _add_run(para, "The course code is ")
    _add_equation(para, "COMP3013")
    _add_run(para, ".")

    path = tmp_path / "math.pptx"
    prs.save(str(path))

    deck = parse_pptx(str(path))
    texts = [b.text for b in deck.pages[0].blocks]
    assert "The course code is COMP3013." in texts, texts


def test_a_paragraph_that_is_only_an_equation_survives(tmp_path):
    """Such a paragraph is empty as far as python-pptx is concerned, so the
    "skip blank paragraphs" filter used to delete it outright."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    _add_equation(box.text_frame.paragraphs[0], "EMC")

    path = tmp_path / "only-math.pptx"
    prs.save(str(path))

    deck = parse_pptx(str(path))
    assert "EMC" in [b.text for b in deck.pages[0].blocks]


def test_equation_blocks_are_marked_for_the_translator(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    para = box.text_frame.paragraphs[0]
    _add_run(para, "Let ")
    _add_equation(para, "R")
    _add_run(para, " be a relation.")

    path = tmp_path / "marked.pptx"
    prs.save(str(path))

    deck = parse_pptx(str(path))
    block = next(b for b in deck.pages[0].blocks if "relation" in b.text)
    assert block.meta["has_math"] is True


def test_unsupported_format_names_the_fix(tmp_path):
    bad = tmp_path / "slides.key"
    bad.write_bytes(b"")
    with pytest.raises(UnsupportedFormat, match="Keynote"):
        parse(str(bad))
