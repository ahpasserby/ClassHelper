"""The formats that need LibreOffice.

`.ppt` is a compound binary document, not a zip of XML, so it is converted once
and read by the normal parser. The behaviour worth pinning down is the failure:
without LibreOffice this has to say what to install, not "unsupported format".
"""

from __future__ import annotations

import pytest

from classhelper.parsers import UnsupportedFormat, legacy, parse


def test_a_ppt_without_libreoffice_explains_the_fix(tmp_path, monkeypatch):
    import classhelper.render as render_module

    monkeypatch.setattr(render_module, "find_soffice", lambda: None)
    path = tmp_path / "old.ppt"
    path.write_bytes(b"\xd0\xcf\x11\xe0")  # a compound document's magic number

    with pytest.raises(UnsupportedFormat) as raised:
        parse(str(path))
    message = str(raised.value)
    assert "LibreOffice" in message
    assert ".pptx" in message


def test_a_converted_ppt_still_reports_where_it_came_from(tmp_path, monkeypatch):
    """The board, the caches and "reveal in Finder" all point at the file the
    user has, not at the copy in the cache directory."""
    from pptx import Presentation

    real = tmp_path / "converted.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Old deck"
    prs.save(str(real))

    source = tmp_path / "old.ppt"
    source.write_bytes(b"\xd0\xcf\x11\xe0")
    monkeypatch.setattr(legacy, "convert", lambda path, to: real)

    deck = parse(str(source))
    assert deck.source_path == str(source)
    assert deck.source_format == "ppt"
    assert deck.pages[0].blocks[0].text == "Old deck"
    assert deck.notes


def test_an_unknown_format_names_the_ones_that_work(tmp_path):
    path = tmp_path / "slides.key"
    path.write_bytes(b"")
    with pytest.raises(UnsupportedFormat) as raised:
        parse(str(path))
    assert ".pptx" in str(raised.value) and ".docx" in str(raised.value)
