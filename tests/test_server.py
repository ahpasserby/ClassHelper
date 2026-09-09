"""The HTTP layer, with a fake model.

These run offline: no API key, no network. What they cover is the plumbing that
broke in practice -- the snapshot a client needs on connect, and the fact that
a correction has to reach the reader.
"""

from __future__ import annotations

import json
import os

import pytest


@pytest.fixture
def live_server(client):
    """A real uvicorn on a real port, sharing the fixtures' patched app.

    Only the event stream needs this; everything else is faster and simpler
    through the in-process client.
    """
    import threading

    import uvicorn

    import classhelper.server.app as app_module
    from classhelper.launcher import free_port, wait_until_ready

    port = free_port()
    config = uvicorn.Config(app_module.app, host="127.0.0.1", port=port,
                            log_level="error")
    server = uvicorn.Server(config)
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    assert wait_until_ready(port, timeout=15), "the test server did not start"
    try:
        yield f"http://127.0.0.1:{port}"
    finally:
        server.should_exit = True
        thread.join(timeout=10)


def test_opening_a_deck_returns_its_structure(client, deck_file):
    body = client.post("/api/open", json={"path": deck_file}).json()
    assert len(body["pages"]) == 3
    assert body["pages"][0]["title"] == "Slide 1"


def test_opening_a_missing_file_is_a_404(client):
    response = client.post("/api/open", json={"path": "/nope/absent.pptx"})
    assert response.status_code == 404


def test_an_unsupported_format_explains_the_fix(client, tmp_path):
    bad = tmp_path / "slides.key"
    bad.write_bytes(b"")
    response = client.post("/api/open", json={"path": str(bad)})
    assert response.status_code == 415
    assert "Keynote" in response.json()["detail"]


def test_reopening_the_same_file_resumes_the_session(client, deck_file):
    first = client.post("/api/open", json={"path": deck_file}).json()
    second = client.post("/api/open", json={"path": deck_file}).json()
    assert first["id"] == second["id"]


def test_the_event_stream_opens_with_a_full_snapshot(live_server, deck_file):
    """Regression: a fully cached deck finishes translating before the browser
    connects, so a stream carrying only future events left every sentence
    showing "translating..." forever.

    Run against a real socket rather than the in-process TestClient: this
    endpoint holds the connection open until the client goes away, and ASGI
    transport has no disconnect to detect, so a streamed request there never
    terminates.
    """
    import httpx

    deck = httpx.post(f"{live_server}/api/open", json={"path": deck_file}).json()

    with httpx.stream(
        "GET", f"{live_server}/api/deck/{deck['id']}/events", timeout=10
    ) as stream:
        event = None
        for line in stream.iter_lines():
            if line.startswith("event:"):
                event = line.removeprefix("event:").strip()
            elif line.startswith("data:"):
                payload = json.loads(line[5:])
                assert event == "sync", f"first event was {event!r}, not a snapshot"
                assert len(payload["pages"]) == 3
                return
    pytest.fail("the stream closed without sending anything")


def test_locking_a_term_reports_the_pages_it_changed(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    response = client.post(
        f"/api/deck/{deck['id']}/glossary",
        json={"term": "relationship set", "translation": "联系集"},
    )
    assert response.status_code == 200
    assert sorted(response.json()["pages_changed"]) == [0, 1, 2]


def test_locking_requires_both_halves(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    response = client.post(
        f"/api/deck/{deck['id']}/glossary", json={"term": " ", "translation": "x"}
    )
    assert response.status_code == 400


def test_an_edit_survives_a_reload(client, deck_file):
    """The user's own wording has to outlast the session, or correcting a
    sentence is busywork."""
    deck = client.post("/api/open", json={"path": deck_file}).json()
    sentence = next(
        s for b in deck["pages"][0]["blocks"] for s in b["sentences"]
    )
    client.put(
        f"/api/deck/{deck['id']}/sentence/{sentence['id']}",
        json={"translation": "我自己写的"},
    )
    fresh = client.get(f"/api/deck/{deck['id']}").json()
    edited = next(
        s for b in fresh["pages"][0]["blocks"] for s in b["sentences"]
        if s["id"] == sentence["id"]
    )
    assert edited["translation"] == "我自己写的"
    assert edited["edited"] is True


def test_asking_streams_an_answer(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    with client.stream(
        "POST", f"/api/deck/{deck['id']}/ask",
        json={"page": 0, "question": "为什么？"},
    ) as stream:
        chunks = [
            json.loads(line[5:])["text"]
            for line in stream.iter_lines()
            if line.startswith("data:") and "[DONE]" not in line
        ]
    assert "".join(chunks) == "答案"


def test_asking_about_a_page_that_does_not_exist_is_a_404(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    response = client.post(
        f"/api/deck/{deck['id']}/ask", json={"page": 99, "question": "?"}
    )
    assert response.status_code == 404


def test_an_unknown_deck_is_a_404(client):
    assert client.get("/api/deck/nosuchdeck").status_code == 404


def test_dropped_files_are_imported_from_their_contents(client, deck_file):
    """A browser never reveals a dropped file's path, but it does hand over the
    bytes -- which is all the parser needed. This is the path drag-and-drop takes."""
    with open(deck_file, "rb") as handle:
        response = client.post(
            "/api/upload",
            files={"files": ("Lec01.pptx", handle.read(),
                             "application/vnd.openxmlformats-officedocument."
                             "presentationml.presentation")},
        )
    body = response.json()
    assert response.status_code == 200
    assert [i["name"] for i in body["added"]] == ["Lec01"]
    assert [i["name"] for i in body["board"]["inbox"]] == ["Lec01"]


def test_importing_does_not_open_anything(client, deck_file):
    """Filing and reading are separate decisions. Dropping a term's worth of
    lectures on the window means "file these", not "open all twelve"."""
    with open(deck_file, "rb") as handle:
        data = handle.read()
    client.post(
        "/api/upload",
        files=[("files", (f"Lec{n}.pptx", data, "application/octet-stream"))
               for n in range(3)],
    )
    import classhelper.server.app as app_module

    assert app_module.store.all() == [], "importing must not start a reader session"


def test_several_files_open_in_one_drop(client, deck_file):
    with open(deck_file, "rb") as handle:
        data = handle.read()
    response = client.post(
        "/api/upload",
        files=[
            ("files", ("Lec01.pptx", data, "application/octet-stream")),
            ("files", ("Lec02.pptx", data, "application/octet-stream")),
        ],
    )
    added = response.json()["added"]
    assert [d["name"] for d in added] == ["Lec01", "Lec02"]
    assert added[0]["id"] != added[1]["id"]


def test_a_folder_of_mixed_files_is_filed_whole(client, deck_file):
    """A course hands out more than decks. Everything is filed; the ones the
    reader cannot open are marked, not refused."""
    with open(deck_file, "rb") as handle:
        data = handle.read()
    response = client.post(
        "/api/upload",
        files=[
            ("files", ("good.pptx", data, "application/octet-stream")),
            ("files", ("syllabus.txt", b"weeks", "text/plain")),
        ],
    )
    body = response.json()
    assert [d["name"] for d in body["added"]] == ["good", "syllabus"]
    assert body["failed"] == []
    assert [d["readable"] for d in body["added"]] == [True, False]


def test_one_bad_file_does_not_sink_the_rest(client, deck_file):
    """Dropping a folder's worth of files should take what it can and say what
    it could not, rather than refusing the whole batch."""
    with open(deck_file, "rb") as handle:
        data = handle.read()
    response = client.post(
        "/api/upload",
        files=[
            ("files", ("good.pptx", data, "application/octet-stream")),
            ("files", ("empty.pdf", b"", "application/pdf")),
        ],
    )
    body = response.json()
    assert [d["name"] for d in body["added"]] == ["good"]
    assert body["failed"][0]["name"] == "empty.pdf"


def test_a_drop_of_only_unusable_files_is_an_error(client):
    response = client.post(
        "/api/upload",
        files={"files": ("empty.pptx", b"", "application/octet-stream")},
    )
    assert response.status_code == 422


def test_opening_something_that_is_not_a_deck_says_so(client, tmp_path):
    """It is on the board on purpose. Trying to read it has to fail in a way
    that names the file and the formats that do work."""
    other = tmp_path / "syllabus.txt"
    other.write_text("weeks", encoding="utf-8")
    response = client.post("/api/open", json={"path": str(other)})
    assert response.status_code == 415
    assert ".pdf" in response.json()["detail"]


def test_dropping_the_same_name_twice_keeps_both(client, deck_file):
    """The library is a directory, so a second file of the same name is a
    second file -- numbered rather than silently overwriting the first."""
    with open(deck_file, "rb") as handle:
        data = handle.read()
    client.post("/api/upload",
                files={"files": ("a.pptx", data, "application/octet-stream")})
    body = client.post(
        "/api/upload", files={"files": ("a.pptx", data, "application/octet-stream")}
    ).json()
    assert sorted(i["name"] for i in body["board"]["inbox"]) == ["a", "a (2)"]


def test_files_chosen_from_the_dialog_are_imported_not_opened(client, deck_file):
    body = client.post("/api/import", json={"paths": [deck_file]}).json()
    assert [i["name"] for i in body["added"]] == ["deck"]
    import classhelper.server.app as app_module

    assert app_module.store.all() == []


def test_importing_a_missing_path_is_reported_per_file(client, deck_file):
    body = client.post(
        "/api/import", json={"paths": [deck_file, "/nope/absent.pptx"]}
    ).json()
    assert len(body["added"]) == 1
    assert body["failed"][0]["name"] == "absent.pptx"


def test_closing_a_deck_forgets_it(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    assert client.delete(f"/api/deck/{deck['id']}").status_code == 200
    assert client.get(f"/api/deck/{deck['id']}").status_code == 404


def test_closing_an_unknown_deck_is_not_an_error(client):
    """The reader closes a deck when its tab goes away, which can happen after
    the deck is already gone. That must not surface as a failure."""
    assert client.delete("/api/deck/nosuchdeck").status_code == 200


def test_two_decks_stay_independent(client, deck_file, tmp_path):
    from pptx import Presentation

    other = Presentation()
    other.slides.add_slide(other.slide_layouts[5]).shapes.title.text = "Other"
    other_path = tmp_path / "other.pptx"
    other.save(str(other_path))

    a = client.post("/api/open", json={"path": deck_file}).json()
    b = client.post("/api/open", json={"path": str(other_path)}).json()
    assert a["id"] != b["id"]
    assert len(a["pages"]) == 3 and len(b["pages"]) == 1
    assert client.get(f"/api/deck/{a['id']}").json()["pages"][0]["title"] == "Slide 1"


# -- the board over HTTP ---------------------------------------------------

@pytest.fixture
def board_client(client):
    """The board endpoints share the `client` fixture's scratch library."""
    return client


def test_the_board_starts_empty(board_client):
    body = board_client.get("/api/board").json()
    assert body["folders"] == [] and body["inbox"] == []
    assert body["root"].endswith("library")


def test_folders_are_real_directories(board_client, tmp_path):
    semester = board_client.post(
        "/api/board/folders", json={"name": "2026 秋季"}
    ).json()["id"]
    board_client.post("/api/board/folders", json={"name": "DMS", "parent": semester})

    assert (tmp_path / "library" / "2026 秋季" / "DMS").is_dir()
    folders = board_client.get("/api/board").json()["folders"]
    assert folders[0]["name"] == "2026 秋季"
    assert [c["name"] for c in folders[0]["children"]] == ["DMS"]
    assert folders[0]["children"][0]["id"] == "2026 秋季/DMS"


def test_a_dropped_file_lands_in_the_inbox_directory(board_client, deck_file, tmp_path):
    with open(deck_file, "rb") as handle:
        board_client.post(
            "/api/upload",
            files={"files": ("Lec01.pptx", handle.read(), "application/octet-stream")},
        )
    body = board_client.get("/api/board").json()
    assert [i["name"] for i in body["inbox"]] == ["Lec01"]
    assert (tmp_path / "library" / "_inbox" / "Lec01.pptx").exists()


def test_filing_a_deck_really_moves_the_file(board_client, deck_file, tmp_path):
    folder = board_client.post("/api/board/folders", json={"name": "DMS"}).json()["id"]
    item = board_client.post(
        "/api/import", json={"paths": [deck_file]}
    ).json()["added"][0]["id"]

    board = board_client.post(
        "/api/board/move", json={"paths": [item], "dest": folder}
    ).json()["board"]
    assert board["inbox"] == []
    assert (tmp_path / "library" / "DMS" / "deck.pptx").exists()
    assert not (tmp_path / "library" / "_inbox" / "deck.pptx").exists()


def test_a_selection_moves_in_one_request(board_client, deck_file, tmp_path):
    """Filing a week of lectures is one gesture. Doing it as N requests leaves
    the board visibly half-moved when one of them fails."""
    from pptx import Presentation

    folder = board_client.post("/api/board/folders", json={"name": "DMS"}).json()["id"]
    sources = []
    for n in range(3):
        path = tmp_path / f"deck{n}.pptx"
        Presentation().save(str(path))
        sources.append(str(path))
    ids = [
        i["id"]
        for i in board_client.post(
            "/api/import", json={"paths": sources}
        ).json()["added"]
    ]

    body = board_client.post(
        "/api/board/move", json={"paths": ids, "dest": folder}
    ).json()
    assert len(body["moved"]) == 3
    assert body["board"]["inbox"] == []
    assert len(body["board"]["folders"][0]["items"]) == 3


def test_moving_a_selection_back_unfiles_it(board_client, deck_file):
    folder = board_client.post("/api/board/folders", json={"name": "DMS"}).json()["id"]
    item = board_client.post(
        "/api/import", json={"paths": [deck_file], "dest": folder}
    ).json()["added"][0]["id"]
    body = board_client.post(
        "/api/board/move", json={"paths": [item], "dest": None}
    ).json()
    assert [i["name"] for i in body["board"]["inbox"]] == ["deck"]


def test_moving_a_stale_path_is_skipped_not_fatal(board_client, deck_file):
    """A path from a board the reader has not refreshed must not take the rest
    of the selection down with it."""
    item = board_client.post(
        "/api/import", json={"paths": [deck_file]}
    ).json()["added"][0]["id"]
    folder = board_client.post("/api/board/folders", json={"name": "X"}).json()["id"]
    body = board_client.post(
        "/api/board/move", json={"paths": [item, "_inbox/goneforever.pptx"],
                                 "dest": folder}
    ).json()
    assert len(body["moved"]) == 1


def test_sibling_courses_can_be_added_one_after_another(board_client):
    """Regression for "once there is one course you cannot add another"."""
    semester = board_client.post(
        "/api/board/folders", json={"name": "S"}
    ).json()["id"]
    for name in ("DMS", "OS", "Compilers"):
        response = board_client.post(
            "/api/board/folders", json={"name": name, "parent": semester}
        )
        assert response.status_code == 200, response.json()
    folders = board_client.get("/api/board").json()["folders"]
    # Directory listings are sorted by name, not by creation order.
    assert sorted(c["name"] for c in folders[0]["children"]) == [
        "Compilers", "DMS", "OS",
    ]


def test_a_folder_cannot_swallow_itself(board_client):
    parent = board_client.post("/api/board/folders", json={"name": "A"}).json()["id"]
    child = board_client.post(
        "/api/board/folders", json={"name": "B", "parent": parent}
    ).json()["id"]
    response = board_client.post(
        "/api/board/move", json={"paths": [parent], "dest": child}
    )
    assert response.status_code == 400


def test_a_path_cannot_escape_the_library(board_client):
    """Folder names arrive from the browser; without the guard `../..` is one."""
    response = board_client.post(
        "/api/board/move", json={"paths": ["../../etc/passwd"], "dest": None}
    )
    assert response.status_code == 400


def test_deleting_a_folder_rescues_the_decks_inside_it(board_client, deck_file):
    folder = board_client.post("/api/board/folders", json={"name": "DMS"}).json()["id"]
    board_client.post("/api/import", json={"paths": [deck_file], "dest": folder})
    body = board_client.post("/api/board/delete-folder", json={"path": folder}).json()
    assert body["returned_to_inbox"] == 1
    assert len(body["board"]["inbox"]) == 1


def test_a_broken_link_is_marked_not_hidden(board_client, tmp_path):
    """Linked imports point at a file the library does not own; saying so on
    the board beats failing only when it is opened."""
    from pptx import Presentation

    original = tmp_path / "linked.pptx"
    Presentation().save(str(original))
    board_client.get("/api/board")  # ensure the library exists
    link = tmp_path / "library" / "_inbox" / "linked.pptx"
    link.parent.mkdir(parents=True, exist_ok=True)
    link.symlink_to(original)
    original.unlink()

    inbox = board_client.get("/api/board").json()["inbox"]
    assert [i["missing"] for i in inbox] == [True]


def test_a_glossary_scope_inherits_from_its_ancestors(board_client):
    semester = board_client.post(
        "/api/board/folders", json={"name": "2026 秋季"}
    ).json()["id"]
    course = board_client.post(
        "/api/board/folders", json={"name": "DMS", "parent": semester}
    ).json()["id"]

    board_client.post(
        "/api/glossary", json={"term": "relationship", "translation": "联系",
                               "scope": semester}
    )
    body = board_client.get(f"/api/glossary?scope={course}").json()
    assert [s["name"] for s in body["chain"]] == ["全局", "2026 秋季", "DMS"]
    assert body["terms"][0]["translation"] == "联系"
    assert body["terms"][0]["inherited"] is True


def test_an_override_is_local_to_its_course(board_client):
    semester = board_client.post(
        "/api/board/folders", json={"name": "S"}
    ).json()["id"]
    dms = board_client.post(
        "/api/board/folders", json={"name": "DMS", "parent": semester}
    ).json()["id"]
    other = board_client.post(
        "/api/board/folders", json={"name": "OS", "parent": semester}
    ).json()["id"]

    board_client.post("/api/glossary",
                      json={"term": "key", "translation": "键", "scope": semester})
    board_client.post("/api/glossary",
                      json={"term": "key", "translation": "码", "scope": dms})

    here = board_client.get(f"/api/glossary?scope={dms}").json()["terms"]
    there = board_client.get(f"/api/glossary?scope={other}").json()["terms"]
    assert [t["translation"] for t in here] == ["码"]
    assert [t["translation"] for t in there] == ["键"]


def test_dropping_an_override_reveals_the_inherited_term(board_client):
    semester = board_client.post("/api/board/folders", json={"name": "S"}).json()["id"]
    course = board_client.post(
        "/api/board/folders", json={"name": "C", "parent": semester}
    ).json()["id"]
    board_client.post("/api/glossary",
                      json={"term": "key", "translation": "键", "scope": semester})
    board_client.post("/api/glossary",
                      json={"term": "key", "translation": "码", "scope": course})
    board_client.delete(f"/api/glossary/key?scope={course}")

    terms = board_client.get(f"/api/glossary?scope={course}").json()["terms"]
    assert [(t["translation"], t["inherited"]) for t in terms] == [("键", True)]


def test_an_unknown_glossary_scope_is_a_404(board_client):
    assert board_client.get("/api/glossary?scope=nope").status_code == 404


# -- settings --------------------------------------------------------------

@pytest.fixture
def settings_client(client, tmp_path, monkeypatch):
    """The API client with the settings file pointed at a scratch path.

    Without this the suite would overwrite the developer's real config, API key
    and all.
    """
    import classhelper.config as config_module
    import classhelper.server.app as app_module

    path = tmp_path / "home" / "config.toml"
    monkeypatch.setattr(config_module, "CONFIG_PATH", path)
    monkeypatch.setattr(app_module, "CONFIG_PATH", path)
    monkeypatch.setattr(app_module, "load_config", config_module.load)
    monkeypatch.setattr(app_module, "save_config", config_module.save)
    monkeypatch.delenv("CLASSHELPER_API_KEY", raising=False)
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    # Saving settings re-points the library, so every write here must name the
    # scratch one -- otherwise the suite relocates the developer's own.
    import classhelper.server.board_api as board_module

    monkeypatch.setattr(
        app_module, "DEFAULT_USER_DATA", tmp_path / "userdata"
    )
    return client, path


def test_settings_report_that_nothing_is_configured_yet(settings_client):
    client, _ = settings_client
    body = client.get("/api/settings").json()
    assert body["configured"] is False
    assert body["api_key_set"] is False
    assert "deepseek" in [p["id"] for p in body["providers"]]


def test_a_first_run_shows_the_defaults_and_writes_nothing(settings_client):
    """A fresh install once came up saying it was configured, with a maximised
    window it had never been asked for -- reading settings must not create them,
    and the page it feeds should show the recommended models, not blank fields.
    """
    client, path = settings_client
    body = client.get("/api/settings").json()
    assert body["configured"] is False
    assert body["model"] and body["ask_model"] and body["base_url"]
    assert body["window_mode"] == "fullscreen"
    assert not path.exists(), "merely looking at the settings wrote a file"


def test_saving_settings_writes_a_private_file(settings_client):
    client, path = settings_client
    response = client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "sk-secret", "target_lang": "zh-CN"},
    )
    assert response.status_code == 200
    assert path.exists()
    assert oct(path.stat().st_mode)[-3:] == "600", "the file holds an API key"
    assert "sk-secret" in path.read_text()


def test_defaults_fill_in_when_switching_provider(settings_client):
    """Switching provider re-seeds the model names; keeping the old provider's
    models pointed at a new service would just 404."""
    client, path = settings_client
    client.put("/api/settings", json={"provider": "deepseek",
                                      "api_key": "sk-fake-0123456789"})
    client.put("/api/settings", json={"provider": "openai",
                                      "api_key": "sk-fake-0123456789"})
    body = client.get("/api/settings").json()
    assert body["base_url"] == "https://api.openai.com/v1"
    assert body["model"] and body["ask_model"]


def test_the_key_is_never_sent_back_in_full(settings_client):
    client, _ = settings_client
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-abcdef123456"})
    body = client.get("/api/settings").json()
    assert body["api_key_set"] is True
    assert "sk-abcdef123456" not in json.dumps(body)


def test_an_untouched_key_field_keeps_the_stored_key(settings_client):
    """The page is never told the key, so it cannot send it back -- changing
    the model must not therefore wipe the key."""
    client, path = settings_client
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-keepme"})
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "__unchanged__",
              "model": "deepseek-chat-v2"},
    )
    assert "sk-keepme" in path.read_text()
    assert "deepseek-chat-v2" in path.read_text()


def test_a_provider_needing_a_key_refuses_to_save_without_one(settings_client):
    client, _ = settings_client
    response = client.put("/api/settings", json={"provider": "deepseek", "api_key": ""})
    assert response.status_code == 400


def test_a_local_provider_saves_without_a_key(settings_client):
    client, path = settings_client
    assert client.put(
        "/api/settings", json={"provider": "ollama", "api_key": ""}
    ).status_code == 200
    assert "11434" in path.read_text()


def test_a_custom_provider_must_be_given_a_url_and_models(settings_client):
    client, _ = settings_client
    response = client.put(
        "/api/settings", json={"provider": "custom", "api_key": "sk-fake-0123456789"}
    )
    assert response.status_code == 400
    assert client.put(
        "/api/settings",
        json={"provider": "custom", "api_key": "sk-fake-0123456789",
              "base_url": "https://example.test/v1", "model": "m", "ask_model": "m2"},
    ).status_code == 200


def test_an_unknown_provider_is_refused(settings_client):
    client, _ = settings_client
    response = client.put("/api/settings", json={"provider": "nope", "api_key": "sk-fake-0123456789"})
    assert response.status_code == 400


def test_a_failed_connection_test_reports_rather_than_raises(settings_client):
    """The page needs the reason shown next to the fields, not an HTTP error."""
    client, _ = settings_client
    response = client.post(
        "/api/settings/test",
        json={"provider": "custom", "api_key": "sk-fake-0123456789",
              "base_url": "http://127.0.0.1:9/v1", "model": "m", "ask_model": "m"},
    )
    assert response.status_code == 200
    assert response.json()["ok"] is False
    assert response.json()["detail"]


def test_a_suspiciously_short_key_is_refused(settings_client):
    """A browser autofilling a password field here once overwrote a working key
    with a single character, and the only copy was gone."""
    client, path = settings_client
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-realkey123"})
    response = client.put("/api/settings", json={"provider": "deepseek", "api_key": "1"})
    assert response.status_code == 400
    assert "sk-realkey123" in path.read_text(), "the stored key must survive a refusal"


def test_saving_keeps_a_backup_of_the_previous_file(settings_client):
    client, path = settings_client
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-first12345"})
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-second1234"})
    backup = path.with_name(path.name + ".bak")
    assert backup.exists()
    assert "sk-first12345" in backup.read_text()
    assert oct(backup.stat().st_mode)[-3:] == "600"


def test_the_key_can_be_revealed_on_request(settings_client):
    """The page never receives the key with the rest of the settings; the eye
    control asks for it separately."""
    client, _ = settings_client
    client.put("/api/settings", json={"provider": "deepseek", "api_key": "sk-visible123"})
    assert "sk-visible123" not in json.dumps(client.get("/api/settings").json())
    assert client.get("/api/settings/key").json()["api_key"] == "sk-visible123"


def test_an_unreachable_provider_returns_an_empty_model_list(settings_client):
    """Not every OpenAI-compatible service implements /models. The field stays
    free text rather than the page breaking."""
    client, _ = settings_client
    body = client.post(
        "/api/settings/models",
        json={"provider": "custom", "api_key": "sk-fake-0123456789", "base_url": "http://127.0.0.1:9/v1",
              "model": "m", "ask_model": "m"},
    ).json()
    assert body["models"] == []
    assert body["detail"]


# -- saved conversations ---------------------------------------------------

def test_a_conversation_is_stored_beside_its_deck(client, board_client, deck_file, tmp_path):
    """The point of storing transcripts in the folder: they are still there
    next week, and they travel with the course."""
    folder = board_client.post("/api/board/folders", json={"name": "DMS"}).json()["id"]
    item = board_client.post(
        "/api/import", json={"paths": [deck_file], "dest": folder}
    ).json()["added"][0]["id"]
    deck = client.post("/api/open", json={"path": item}).json()

    chat = client.post(f"/api/deck/{deck['id']}/chats").json()
    client.post(
        f"/api/deck/{deck['id']}/chats/{chat['id']}/messages",
        json={"parentId": None,
              "message": {"role": "user",
                          "content": [{"type": "text", "text": "这是什么意思？"}]}},
    )

    stored = tmp_path / "library" / "DMS" / ".classhelper" / "chats" / "deck.json"
    assert stored.exists()
    assert "这是什么意思？" in stored.read_text(encoding="utf-8")


def test_a_conversation_is_named_after_its_first_question(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    chat = client.post(f"/api/deck/{deck['id']}/chats").json()
    assert chat["title"] == "新对话"
    client.post(
        f"/api/deck/{deck['id']}/chats/{chat['id']}/messages",
        json={"parentId": None,
              "message": {"role": "user",
                          "content": [{"type": "text", "text": "agree on 是什么意思"}]}},
    )
    listed = client.get(f"/api/deck/{deck['id']}/chats").json()["chats"]
    assert listed[0]["title"] == "agree on 是什么意思"


def test_conversations_survive_reopening_the_deck(client, deck_file):
    """Reopening must find the transcript, which is the whole feature."""
    deck = client.post("/api/open", json={"path": deck_file}).json()
    chat = client.post(f"/api/deck/{deck['id']}/chats").json()
    client.post(
        f"/api/deck/{deck['id']}/chats/{chat['id']}/messages",
        json={"parentId": None,
              "message": {"role": "user", "content": [{"type": "text", "text": "hi"}]}},
    )
    client.delete(f"/api/deck/{deck['id']}")

    again = client.post("/api/open", json={"path": deck_file}).json()
    body = client.get(f"/api/deck/{again['id']}/chats/{chat['id']}").json()
    assert len(body["messages"]) == 1


def test_conversations_can_be_deleted(client, deck_file):
    deck = client.post("/api/open", json={"path": deck_file}).json()
    keep = client.post(f"/api/deck/{deck['id']}/chats").json()["id"]
    drop = client.post(f"/api/deck/{deck['id']}/chats").json()["id"]
    client.delete(f"/api/deck/{deck['id']}/chats/{drop}")
    remaining = [c["id"] for c in client.get(f"/api/deck/{deck['id']}/chats").json()["chats"]]
    assert remaining == [keep]


def test_appending_to_a_deleted_conversation_does_not_lose_the_message(client, deck_file):
    """The reader appends as the exchange happens; a transcript removed in
    another window must not swallow the answer being written right now."""
    deck = client.post("/api/open", json={"path": deck_file}).json()
    chat = client.post(f"/api/deck/{deck['id']}/chats").json()["id"]
    client.delete(f"/api/deck/{deck['id']}/chats/{chat}")
    client.post(
        f"/api/deck/{deck['id']}/chats/{chat}/messages",
        json={"parentId": None,
              "message": {"role": "user", "content": [{"type": "text", "text": "x"}]}},
    )
    assert len(client.get(f"/api/deck/{deck['id']}/chats/{chat}").json()["messages"]) == 1


def test_the_two_user_directories_derive_from_one_setting(settings_client, tmp_path):
    """One place to point at a synced folder, split into the material and the
    program's own state."""
    client, _ = settings_client
    chosen = tmp_path / "somewhere"
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "sk-fake-0123456789",
              "user_data_path": str(chosen)},
    )
    body = client.get("/api/settings").json()
    assert body["user_data_path"] == str(chosen)
    assert body["library_path"] == str(chosen / "课板")
    assert body["state_path"] == str(chosen / "用户配置")


def test_the_window_mode_is_remembered(settings_client):
    client, path = settings_client
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "sk-fake-0123456789",
              "window_mode": "maximized"},
    )
    assert 'mode = "maximized"' in path.read_text()
    assert client.get("/api/settings").json()["window_mode"] == "maximized"


def test_saving_settings_does_not_relocate_the_library(settings_client, tmp_path):
    """Regression: a save that did not mention the library reset it to the
    default path, and the next import went into the wrong tree."""
    client, _ = settings_client
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "sk-fake-0123456789",
              "user_data_path": str(tmp_path / "chosen")},
    )
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "__unchanged__",
              "target_lang": "en"},
    )
    assert client.get("/api/settings").json()["user_data_path"] == str(
        tmp_path / "chosen"
    )


def test_an_unmentioned_field_is_not_reset_to_the_provider_default(settings_client):
    """The bug behind "my settings never save": a save that mentioned only the
    language also rewrote both model names back to the defaults."""
    client, _ = settings_client
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "sk-fake-0123456789",
              "model": "my-fast-model", "ask_model": "my-smart-model"},
    )
    client.put(
        "/api/settings",
        json={"provider": "deepseek", "api_key": "__unchanged__",
              "target_lang": "ja"},
    )
    body = client.get("/api/settings").json()
    assert body["model"] == "my-fast-model"
    assert body["ask_model"] == "my-smart-model"
    assert body["target_lang"] == "ja"


# -- the contract the desktop shell depends on -----------------------------

def test_the_cli_can_serve_headlessly_on_a_chosen_port(tmp_path, monkeypatch):
    """The Electron shell starts the server with exactly this invocation and
    then polls /api/status. If the flags or that endpoint change, the desktop
    app stops launching -- with no error a user could act on."""
    import subprocess
    import sys
    import urllib.request

    from classhelper.launcher import free_port, wait_until_ready

    port = free_port()
    process = subprocess.Popen(
        [sys.executable, "-m", "classhelper.cli", "serve", "--port", str(port)],
        env={**os.environ, "CLASSHELPER_HOME": str(tmp_path), "NO_PROXY": "*"},
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    try:
        assert wait_until_ready(port, timeout=40), "the server never answered"
        with urllib.request.urlopen(
            f"http://127.0.0.1:{port}/api/status", timeout=5
        ) as response:
            assert response.status == 200
    finally:
        process.terminate()
        process.wait(timeout=10)


def test_the_frozen_entry_point_accepts_the_shell_arguments(tmp_path):
    """PyInstaller freezes `classhelper.cli:main`. Parsing has to accept the
    exact argument list the desktop shell passes, or the packaged app dies at
    launch with an argparse usage message nobody sees."""
    import argparse
    from unittest.mock import patch

    from classhelper import cli

    with patch("classhelper.launcher.serve", return_value=0) as serve:
        assert cli.main(["serve", "--port", "54321"]) == 0
    serve.assert_called_once_with(54321)

    with pytest.raises(SystemExit):
        cli.main(["serve", "--nonsense"])
    assert argparse  # the parser is what enforces the contract above
